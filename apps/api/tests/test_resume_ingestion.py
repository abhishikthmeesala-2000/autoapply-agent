from __future__ import annotations

from io import BytesIO

import pytest
from docx import Document
from fastapi.testclient import TestClient
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas

from app.db.models import Profile, User
from app.main import app, get_db_session


class FakeVectorStore:
    def __init__(self) -> None:
        self.calls: list[dict[str, object]] = []

    def upsert_resume_evidence(
        self,
        *,
        profile_id: str,
        master_resume_id: str,
        evidence: list[tuple[str, str, str]],
    ) -> None:
        self.calls.append(
            {
                "profile_id": profile_id,
                "master_resume_id": master_resume_id,
                "evidence": evidence,
            }
        )


def _build_docx_resume() -> bytes:
    document = Document()
    document.add_paragraph("SUMMARY")
    document.add_paragraph("Product engineer with 7 years of experience.")
    document.add_paragraph("SKILLS")
    document.add_paragraph("- Python")
    document.add_paragraph("- FastAPI")
    document.add_paragraph("EXPERIENCE")
    document.add_paragraph("Acme Corp | Senior Engineer | 2022-01 to Present")
    document.add_paragraph("- Built internal automation tools.")
    document.add_paragraph("PROJECTS")
    document.add_paragraph("- Job automation platform for local workflows.")
    document.add_paragraph("EDUCATION")
    document.add_paragraph("B.S. Computer Science, State University, 2018")
    document.add_paragraph("CERTIFICATIONS")
    document.add_paragraph("- AWS Certified Developer Associate")

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _build_pdf_resume() -> bytes:
    buffer = BytesIO()
    pdf = canvas.Canvas(buffer, pagesize=letter)
    y = 750
    for line in [
        "SUMMARY",
        "Backend engineer with strong API and data skills.",
        "SKILLS",
        "- Python",
        "- SQLAlchemy",
        "EXPERIENCE",
        "Beta LLC | Engineer | 2020-2024",
        "- Built data pipelines.",
        "PROJECTS",
        "- Resume parser with evidence mapping.",
        "EDUCATION",
        "B.S. Software Engineering, Metro University, 2017",
        "CERTIFICATIONS",
        "- CKAD",
    ]:
        pdf.drawString(72, y, line)
        y -= 18
    pdf.showPage()
    pdf.save()
    return buffer.getvalue()


@pytest.fixture()
def api_client(db_session):
    fake_vector_store = FakeVectorStore()

    def override_get_db_session():
        yield db_session

    app.dependency_overrides[get_db_session] = override_get_db_session
    app.state.resume_vector_store = fake_vector_store

    from app import main as main_module

    original_store = main_module.vector_store
    main_module.vector_store = fake_vector_store

    client = TestClient(app)
    try:
        yield client, fake_vector_store
    finally:
        app.dependency_overrides.clear()
        main_module.vector_store = original_store


def _create_profile(db_session) -> Profile:
    user = User(email="resume-test@example.com")
    db_session.add(user)
    db_session.flush()
    profile = Profile(user_id=user.id, name="Target", location="Remote")
    db_session.add(profile)
    db_session.commit()
    return profile


def test_resume_upload_docx_parses_and_persists(api_client, db_session) -> None:
    client, fake_vector_store = api_client
    profile = _create_profile(db_session)

    response = client.post(
        f"/profiles/{profile.id}/resume/upload",
        files={
            "upload": (
                "resume.docx",
                _build_docx_resume(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )

    assert response.status_code == 200
    body = response.json()
    assert body["status"] == "parsed"
    assert body["sections"] == 6
    assert body["evidence_count"] >= 6
    assert len(fake_vector_store.calls) == 1

    master_response = client.get(f"/profiles/{profile.id}/resume/master")
    assert master_response.status_code == 200
    master_body = master_response.json()
    assert master_body["profile_id"] == profile.id
    assert len(master_body["parsed_json"]["sections"]) == 6
    assert all(
        item["evidence_id"]
        for section in master_body["parsed_json"]["sections"]
        for item in section["items"]
    )

    evidence_response = client.get(f"/profiles/{profile.id}/resume/evidence")
    assert evidence_response.status_code == 200
    evidence_body = evidence_response.json()
    assert len(evidence_body) == body["evidence_count"]
    assert {item["profile_id"] for item in evidence_body} == {profile.id}


def test_resume_upload_pdf_parses_and_persists(api_client, db_session) -> None:
    client, _fake_vector_store = api_client
    profile = _create_profile(db_session)

    response = client.post(
        f"/profiles/{profile.id}/resume/upload",
        files={"upload": ("resume.pdf", _build_pdf_resume(), "application/pdf")},
    )

    assert response.status_code == 200
    assert response.json()["status"] == "parsed"


def test_resume_upload_rejects_invalid_files(api_client, db_session) -> None:
    client, _ = api_client
    profile = _create_profile(db_session)

    response = client.post(
        f"/profiles/{profile.id}/resume/upload",
        files={"upload": ("resume.txt", b"not a supported resume", "text/plain")},
    )

    assert response.status_code == 400
    assert "Only PDF and DOCX are supported" in response.json()["detail"]
