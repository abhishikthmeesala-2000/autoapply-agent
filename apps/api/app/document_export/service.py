from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
from typing import Iterable
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from pypdf import PdfReader
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer
from sqlalchemy.orm import Session

from app.db.models import Profile, ResumeVersion
from app.resume_tailoring.schemas import TailoredResumePayload


class DocumentExportError(RuntimeError):
    pass


@dataclass(frozen=True)
class ExportedResumePackage:
    archive_name: str
    docx_name: str
    pdf_name: str
    archive_bytes: bytes


def _fetch_profile(db: Session, profile_id: str) -> Profile:
    profile = db.get(Profile, profile_id)
    if profile is None:
        raise LookupError("Profile not found.")
    return profile


def _fetch_resume_version(db: Session, profile_id: str, resume_version_id: str) -> ResumeVersion:
    resume_version = db.get(ResumeVersion, resume_version_id)
    if resume_version is None or resume_version.profile_id != profile_id:
        raise LookupError("Resume version not found for this profile.")
    return resume_version


def _load_tailored_payload(resume_version: ResumeVersion) -> TailoredResumePayload:
    try:
        return TailoredResumePayload.model_validate(resume_version.tailored_json)
    except Exception as exc:  # pragma: no cover - defensive validation path
        raise DocumentExportError("Stored tailored resume payload is invalid.") from exc


def _base_file_stem(resume_version: ResumeVersion) -> str:
    stem = Path(resume_version.file_name).name
    stem = Path(stem).stem.strip()
    return stem or f"resume-version-{resume_version.id}"


def _paragraph(document: Document, text: str, *, bold: bool = False, level: int = 0) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0
    if level:
        paragraph.paragraph_format.left_indent = Pt(18 * level)
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = Pt(10.5)


def _build_docx_bytes(profile: Profile, payload: TailoredResumePayload) -> bytes:
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.6)
    section.bottom_margin = Inches(0.6)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(6)
    title_run = title.add_run(profile.name)
    title_run.bold = True
    title_run.font.size = Pt(15)

    meta = document.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.paragraph_format.space_after = Pt(12)
    meta_run = meta.add_run(f"Tailored for {payload.job_id}")
    meta_run.italic = True
    meta_run.font.size = Pt(9)

    for section_payload in payload.sections:
        heading = document.add_paragraph()
        heading.paragraph_format.space_before = Pt(8)
        heading.paragraph_format.space_after = Pt(4)
        heading_run = heading.add_run(section_payload.heading.upper())
        heading_run.bold = True
        heading_run.font.size = Pt(11)

        for bullet in section_payload.bullets:
            _paragraph(document, f"- {bullet.text}", level=0)

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _pdf_story(payload: TailoredResumePayload, profile: Profile) -> Iterable[object]:
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "ResumeTitle",
        parent=styles["Title"],
        fontName="Helvetica-Bold",
        fontSize=16,
        leading=19,
        alignment=TA_LEFT,
        spaceAfter=8,
    )
    meta_style = ParagraphStyle(
        "ResumeMeta",
        parent=styles["BodyText"],
        fontName="Helvetica-Oblique",
        fontSize=9,
        leading=11,
        spaceAfter=14,
    )
    heading_style = ParagraphStyle(
        "SectionHeading",
        parent=styles["Heading2"],
        fontName="Helvetica-Bold",
        fontSize=11,
        leading=13,
        textColor="#111111",
        spaceBefore=8,
        spaceAfter=4,
    )
    body_style = ParagraphStyle(
        "BulletBody",
        parent=styles["BodyText"],
        fontName="Helvetica",
        fontSize=10,
        leading=12,
        leftIndent=12,
        firstLineIndent=-9,
        spaceAfter=2,
    )

    yield Paragraph(profile.name, title_style)
    yield Paragraph(f"Tailored for job {payload.job_id}", meta_style)

    for section_payload in payload.sections:
        yield Paragraph(section_payload.heading.upper(), heading_style)
        for bullet in section_payload.bullets:
            yield Paragraph(f"- {bullet.text}", body_style)
        yield Spacer(1, 4)


def _build_pdf_bytes(profile: Profile, payload: TailoredResumePayload) -> bytes:
    buffer = BytesIO()
    document = SimpleDocTemplate(
        buffer,
        pagesize=letter,
        leftMargin=0.7 * inch,
        rightMargin=0.7 * inch,
        topMargin=0.6 * inch,
        bottomMargin=0.6 * inch,
        title=profile.name,
        author="CareerOS AI",
        subject="Tailored resume export",
    )
    document.build(list(_pdf_story(payload, profile)))
    return buffer.getvalue()


def _validate_docx_bytes(docx_bytes: bytes) -> None:
    try:
        reader = BytesIO(docx_bytes)
        Document(reader)
    except Exception as exc:  # pragma: no cover - defensive validation path
        raise DocumentExportError("Generated DOCX could not be reopened.") from exc


def _validate_pdf_bytes(pdf_bytes: bytes) -> None:
    try:
        reader = PdfReader(BytesIO(pdf_bytes))
    except Exception as exc:  # pragma: no cover - defensive validation path
        raise DocumentExportError("Generated PDF could not be reopened.") from exc

    if not reader.pages:
        raise DocumentExportError("Generated PDF is empty.")


def export_resume_version_package(
    *,
    db: Session,
    profile_id: str,
    resume_version_id: str,
) -> ExportedResumePackage:
    profile = _fetch_profile(db, profile_id)
    resume_version = _fetch_resume_version(db, profile_id, resume_version_id)
    payload = _load_tailored_payload(resume_version)

    if payload.profile_id != profile_id:
        raise LookupError("Tailored resume payload does not belong to this profile.")

    docx_bytes = _build_docx_bytes(profile, payload)
    pdf_bytes = _build_pdf_bytes(profile, payload)
    _validate_docx_bytes(docx_bytes)
    _validate_pdf_bytes(pdf_bytes)

    base_stem = _base_file_stem(resume_version)
    docx_name = f"{base_stem}.docx"
    pdf_name = f"{base_stem}.pdf"
    archive_name = f"{base_stem}.zip"

    archive_buffer = BytesIO()
    with ZipFile(archive_buffer, mode="w", compression=ZIP_DEFLATED) as archive:
        archive.writestr(docx_name, docx_bytes)
        archive.writestr(pdf_name, pdf_bytes)

    return ExportedResumePackage(
        archive_name=archive_name,
        docx_name=docx_name,
        pdf_name=pdf_name,
        archive_bytes=archive_buffer.getvalue(),
    )
